#!/usr/bin/env python3
"""Validación del módulo de peces (CPUE, relación longitud–peso, K de Fulton, Kn de Le Cren, IBI tipo Karr):
página real de Kenti en Chromium vs. implementación independiente (numpy/scipy). Uso: val_peces.py <n_escenarios> <semilla>"""
import subprocess, time, os, sys, tempfile, json, shutil, math
import numpy as np
from scipy import stats
from playwright.sync_api import sync_playwright
EXE = os.environ.get("KENTI_EXE", "../out/Kenti Limnología/kenti-linux")
NSC, SEED = int(sys.argv[1]), int(sys.argv[2])
ORIG = ["nativa", "exotica"]; TOL = ["intolerante", "intermedia", "tolerante"]; GRE = ["omnivoro", "insectivoro", "piscivoro", "herbivoro"]; HAB = ["bentonica", "columna"]
LAB = {"nativa": "Nativa", "exotica": "Exótica", "intolerante": "Intolerante", "intermedia": "Intermedia", "tolerante": "Tolerante", "omnivoro": "Omnívoro",
       "insectivoro": "Insectívoro", "piscivoro": "Piscívoro", "herbivoro": "Herbívoro", "bentonica": "Bentónica", "columna": "Columna de agua"}
METR = [("S", 1), ("Snat", 1), ("Sint", 1), ("Sben", 1), ("pNat", 1), ("pExo", -1), ("pTol", -1), ("pOmn", -1), ("pIns", 1), ("pPis", 1), ("N", 1), ("pAnom", -1)]

def simular(rng, k):
    ns, nsp = int(rng.integers(3, 8)), int(rng.integers(5, 13))
    C = rng.poisson(rng.lognormal(1.5, 1.0, size=(nsp, 1)) * rng.random((nsp, ns)) * 3) * (rng.random((nsp, ns)) < 0.75)
    for j in range(ns):
        if C[:, j].sum() == 0: C[rng.integers(nsp), j] = int(rng.integers(1, 5))
    C = C[C.sum(1) > 0]; nsp = C.shape[0]
    esf = rng.uniform(1, 12, size=ns).round(2)
    tr = [{} for _ in range(nsp)]
    for i in range(nsp):
        for key, opts in (("origen", ORIG), ("tol", TOL), ("gremio", GRE), ("hab", HAB)):
            if rng.random() < 0.88: tr[i][key] = str(rng.choice(opts))
    mm = (k % 2 == 1); ind = []
    for i in range(nsp):
        a = rng.uniform(0.004, 0.02); b = rng.uniform(2.6, 3.4); Lm = rng.uniform(5, 35)
        for j in range(ns):
            n = int(min(C[i, j], rng.integers(0, 22)))
            for _ in range(n):
                L = round(float(np.clip(rng.lognormal(math.log(Lm), 0.35), 2, 90)), 2); W = round(a * L ** b * math.exp(rng.normal(0, 0.12)), 4)
                ind.append((j, i, L, W, str(rng.choice(["M", "H", "J"])), "Sí" if rng.random() < 0.06 else ""))
    return C, esf, tr, ind, mm

def correr(C, esf, tr, ind, mm):
    ns, nsp = C.shape[1], C.shape[0]
    sit = [f"S{j+1}" for j in range(ns)]; spn = [f"Pez {chr(97+i)}{i}" for i in range(nsp)]
    tsv = "Taxón\t" + "\t".join(sit) + "\n" + "\n".join(spn[i] + "\t" + "\t".join(str(int(x)) for x in C[i]) for i in range(nsp))
    f = (lambda x: f"{x*10:.1f}".replace(".", ",")) if mm else (lambda x: f"{x:.2f}".replace(".", ","))
    itsv = "Sitio\tEspecie\t" + ("Longitud (mm)" if mm else "Longitud (cm)") + "\tPeso (g)\tSexo\tAnomalías\n" + "\n".join(f"{sit[j]}\t{spn[i]}\t{f(L)}\t{str(W).replace(".", ",")}\t{s}\t{a}" for j, i, L, W, s, a in ind)
    rtsv = "Especie\tOrigen\tTolerancia\tGremio trófico\tHábitat\n" + "\n".join(spn[i] + "\t" + "\t".join(LAB.get(tr[i].get(k, ""), "") for k in ("origen", "tol", "gremio", "hab")) for i in range(nsp))
    etsv = "\n".join(f"{sit[j]}\t{esf[j]}" for j in range(ns))
    docs = tempfile.mkdtemp(prefix="kp-"); env = dict(os.environ, KENTI_DOCS=docs, KENTI_NO_BROWSER="1")
    srv = subprocess.Popen([EXE], env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL); time.sleep(1.8); errs = []
    try:
        with sync_playwright() as p:
            b = p.chromium.launch(); pg = b.new_page(viewport={"width": 1400, "height": 900}, locale="es-AR")
            pg.on("pageerror", lambda e: errs.append(str(e)))
            pg.goto("http://127.0.0.1:47860/g/peces/"); pg.wait_for_timeout(1500)
            paste = """([id,txt])=>{const el=document.getElementById(id);el.focus();const dt=new DataTransfer();dt.setData('text/plain',txt);el.dispatchEvent(new ClipboardEvent('paste',{clipboardData:dt,bubbles:true,cancelable:true}));}"""
            pg.evaluate(paste, ["g-0-0", tsv]); pg.wait_for_timeout(1200)
            pg.click("#tab-datos"); pg.wait_for_timeout(300)
            pg.fill("#esf-unit", "horas-red"); pg.wait_for_timeout(300)
            pg.evaluate(paste, ["e-0", etsv]); pg.wait_for_timeout(900)
            for tgt, txt in (("ind", itsv), ("rasgos", rtsv)):
                pg.evaluate("t=>window.openLoadDialog(t)", tgt); pg.wait_for_timeout(300)
                pg.fill("#load-area", txt); pg.click("#load-ok"); pg.wait_for_timeout(1200)
            pg.click("#tab-ind"); pg.wait_for_timeout(800); pg.click("#tab-ibi"); pg.wait_for_timeout(800)
            pg.click("#ibi-propose"); pg.wait_for_timeout(1000)
            pg.click("#tab-datos"); pg.wait_for_timeout(1200); pg.click("#btn-export"); pg.wait_for_timeout(5000)
            b.close()
        import glob
        f = glob.glob(docs + "/**/*.xlsx", recursive=True)
        res = leer(f[0], sit, spn) if f else None
        return res, errs, sit, spn
    finally:
        srv.terminate(); time.sleep(0.3); shutil.rmtree(docs, ignore_errors=True)


import openpyxl
IDS = {"S": "S", "S nat": "Snat", "S intol": "Sint", "S bent": "Sben", "% nat": "pNat", "% exót": "pExo", "% tol": "pTol", "% omn": "pOmn", "% ins": "pIns", "% pisc": "pPis", "N": "N", "% DELT": "pAnom"}
NOMBRES = {"Riqueza total de especies": "S", "Riqueza de especies nativas": "Snat", "Riqueza de especies intolerantes": "Sint", "Riqueza de especies bentónicas": "Sben",
           "% de individuos de especies nativas": "pNat", "% de individuos de especies exóticas": "pExo", "% de individuos de especies tolerantes": "pTol", "% de individuos omnívoros": "pOmn",
           "% de individuos insectívoros / invertívoros": "pIns", "% de individuos piscívoros": "pPis", "Abundancia total (N o CPUE del sitio)": "N", "% de individuos con anomalías (DELT)": "pAnom"}
def leer(path, sit, spn):
    wb = openpyxl.load_workbook(path, data_only=True)
    R6 = [list(r) for r in wb.worksheets[5].iter_rows(values_only=True)]; R7 = [list(r) for r in wb.worksheets[6].iter_rows(values_only=True)]
    out = {}
    h = next(i for i, r in enumerate(R6) if r[0] == "Sitio" and r[1] == "Esfuerzo")
    out["cpue"] = [(r[0], r[3]) for r in R6[h + 1:h + 1 + len(sit)]]
    h = next(i for i, r in enumerate(R6) if r[0] == "Especie" and r[1] == "n" and r[2] == "a")
    fits = {}
    for r in R6[h + 1:]:
        if r[0] is None or r[1] is None: break
        fits[r[0]] = dict(n=r[1], a=r[2], b=r[3], lo=r[4], hi=r[5], se=r[6], r2=r[7], t=r[8], p=r[9])
    out["fits"] = fits
    h = next(i for i, r in enumerate(R6) if r[0] == "Especie" and r[1] == "Índice")
    head = R6[h]; cols = {head[j]: j for j in range(3, len(head)) if head[j] and head[j] != "± desvío" and head[j] != "n"}
    cond = {}
    for r in R6[h + 1:]:
        if r[1] not in ("K de Fulton", "Kn de Le Cren"): break
        cond[(r[0], r[1])] = {c: dict(m=r[j], sd=r[j + 1], n=r[j + 2]) for c, j in cols.items()}
    out["cond"] = cond
    h = next(i for i, r in enumerate(R7) if r[0] == "Sitio" and r[2] == "S")
    head = R7[h]; ibi = []
    for r in R7[h + 1:h + 1 + len(sit)]:
        m = {}
        for j in range(2, len(head)):
            if head[j] in IDS and head[j + 1] == "pts": m[IDS[head[j]]] = dict(v=r[j], score=r[j + 1])
        jb = head.index("IBI"); ibi.append(dict(s=r[0], m=m, sum=r[jb], max=r[jb + 1], k=r[jb + 2], pct=r[jb + 3], cls=r[jb + 4]))
    out["ibi"] = ibi
    lim = {}
    for r in R7:
        if r[0] in NOMBRES and r[2] in ("sí", "no"): lim[NOMBRES[r[0]]] = dict(on=r[2] == "sí", t5=r[3], t3=r[4])
    out["metrics"] = lim
    h = next((i for i, r in enumerate(R7) if r[0] and str(r[0]).startswith("Clase")), None)
    return out

def fmax(d, k, v): d[k] = max(d.get(k, 0.0), float(v))
def rpre(v, d):  # redondeo de la propuesta de límites (idéntico a la regla publicada)
    return max(1, math.ceil(v - 1e-9)) if d == 0 else float("%.3g" % v)
def comparar(C, esf, tr, ind, mm, res, sit, spn):
    ns, nsp = C.shape[1], C.shape[0]; D = {}; flags = {}
    by = {}
    for j, i, L, W, s, a in ind: by.setdefault(i, []).append((j, L, W))
    fits = res["fits"]
    keyof = {i: (spn[i] if spn[i] in fits else None) for i in range(nsp)}
    ref_fit = {}
    for i in range(nsp):
        P = by.get(i, [])
        if len(P) < 5: 
            flags["lw_sin_ajuste_ok"] = flags.get("lw_sin_ajuste_ok", True) and (keyof[i] is None); continue
        x = np.log10([L for _, L, _ in P]); y = np.log10([W for _, _, W in P]); n = len(P)
        lr = stats.linregress(x, y); b = lr.slope; se = lr.stderr; t = (b - 3) / se; p = 2 * stats.t.sf(abs(t), n - 2); tc = stats.t.ppf(0.975, n - 2)
        a = 10 ** lr.intercept; ref_fit[i] = (a, b)
        f = fits[keyof[i]]
        for k, ref in (("a", a), ("b", b), ("se", se), ("lo", b - tc * se), ("hi", b + tc * se), ("r2", lr.rvalue ** 2), ("t", t)):
            fmax(D, "lw_" + k, abs(f[k] - ref) / max(1, abs(ref)))
        fmax(D, "lw_p", abs(f["p"] - p))
    cond = res["cond"]
    for i in range(nsp):
        P = by.get(i, [])
        if not P: continue
        cK = cond.get((spn[i], "K de Fulton")); cKn = cond.get((spn[i], "Kn de Le Cren"))
        if cK is None: continue
        for j in range(ns):
            v = [(L, W) for jj, L, W in P if jj == j]
            if not v: continue
            Ks = np.array([100 * W / L ** 3 for L, W in v]); bs = cK[sit[j]]
            fmax(D, "K_n_ok", 0 if bs["n"] == len(v) else 1)
            fmax(D, "K_media", abs(bs["m"] - Ks.mean()) / Ks.mean())
            if len(v) > 1: fmax(D, "K_sd", abs(bs["sd"] - Ks.std(ddof=1)) / Ks.std(ddof=1))
            if i in ref_fit and cKn:
                a, b = ref_fit[i]; Kn = np.array([W / (a * L ** b) for L, W in v]); fmax(D, "Kn_media", abs(cKn[sit[j]]["m"] - Kn.mean()) / Kn.mean())
    for (s, cp), j in zip(res["cpue"], range(ns)): fmax(D, "cpue", abs(cp - C[:, j].sum() / esf[j]))
    # IBI
    vals = {m: [] for m, _ in METR}
    for j in range(ns):
        pres = [i for i in range(nsp) if C[i, j] > 0]; N = C[:, j].sum()
        def pct(key, val):
            wt = [i for i in pres if tr[i].get(key)]; Nt = sum(C[i, j] for i in wt)
            return None if not wt or not Nt else 100 * sum(C[i, j] for i in wt if tr[i][key] == val) / Nt
        def cnt(key, val):
            wt = [i for i in pres if tr[i].get(key)]; return None if not wt else sum(1 for i in wt if tr[i][key] == val)
        an = [r for r in ind if r[0] == j]
        v = {"S": len(pres), "Snat": cnt("origen", "nativa"), "Sint": cnt("tol", "intolerante"), "Sben": cnt("hab", "bentonica"), "pNat": pct("origen", "nativa"), "pExo": pct("origen", "exotica"),
             "pTol": pct("tol", "tolerante"), "pOmn": pct("gremio", "omnivoro"), "pIns": pct("gremio", "insectivoro"), "pPis": pct("gremio", "piscivoro"), "N": N / esf[j],
             "pAnom": (100 * sum(1 for r in an if r[5]) / len(an)) if an else None}
        for m in vals: vals[m].append(v[m])
    thr = {}
    for m, dr in METR:
        obs = [x for x in vals[m] if x is not None]
        if not obs: continue
        if m in ("S", "Snat", "Sint", "Sben", "N"):
            mx = max(obs)
            if not mx > 0: continue
            d = 0 if m != "N" else -1; thr[m] = (rpre(2 * mx / 3, d), min(rpre(mx / 3, d), rpre(2 * mx / 3, d)))
        else: thr[m] = {"pNat": (80, 50), "pExo": (5, 20), "pTol": (5, 20), "pOmn": (20, 45), "pIns": (45, 20), "pPis": (5, 1), "pAnom": (2, 5)}[m]
    ok_thr = True
    for m, (t5, t3) in thr.items():
        c = res["metrics"].get(m, {}); ok_thr &= bool(c.get("on")) and abs(c["t5"] - t5) < 1e-9 and abs(c["t3"] - t3) < 1e-9
    flags["ibi_limites_ok"] = ok_thr
    nok = 0; tot = 0; vmax = 0.0
    for j, r in enumerate(res["ibi"]):
        s = 0; k = 0
        for m, dr in METR:
            got = r["m"][m]["v"]; ref = vals[m][j]
            tot += 1
            if (got is None) != (ref is None): continue
            if ref is not None: vmax = max(vmax, abs(got - ref))
            nok += 1
            if ref is not None and m in thr:
                t5, t3 = thr[m]; sc = (5 if ref >= t5 else 3 if ref >= t3 else 1) if dr > 0 else (5 if ref <= t5 else 3 if ref <= t3 else 1)
                s += sc; k += 1
                if r["m"][m]["score"] != sc: flags["ibi_puntaje_ok"] = False
        flags.setdefault("ibi_puntaje_ok", True)
        flags["ibi_suma_ok"] = flags.get("ibi_suma_ok", True) and ((r["sum"] == s) if k else r["sum"] is None)
        if k:
            pc = 100 * s / (5 * k); fmax(D, "ibi_pct", abs(r["pct"] - pc))
            cls = "Excelente" if pc >= 90 else "Buena" if pc >= 75 else "Regular" if pc >= 60 else "Pobre" if pc >= 45 else "Muy pobre"
            flags["ibi_clase_ok"] = flags.get("ibi_clase_ok", True) and (r["cls"] == cls)
    D["ibi_metricas_validas"] = f"{nok}/{tot}"; D["ibi_metrica_max_dif"] = vmax
    return D, flags

if __name__ == "__main__":
    rng = np.random.default_rng(SEED); out = []
    for k in range(NSC):
        C, esf, tr, ind, mm = simular(rng, k)
        try:
            res, errs, sit, spn = correr(C, esf, tr, ind, mm); D, F = comparar(C, esf, tr, ind, mm, res, sit, spn)
            r = dict(k=k, ns=C.shape[1], nsp=C.shape[0], ninds=len(ind), mm=mm, errs=errs, D=D, F=F)
        except Exception as e:
            import traceback; traceback.print_exc(); r = dict(k=k, error=repr(e)[:200])
        out.append(r); print(json.dumps(r, ensure_ascii=False, default=str), flush=True)
    json.dump(out, open(f"./res_peces_{SEED}.json", "w"), ensure_ascii=False, default=str)
